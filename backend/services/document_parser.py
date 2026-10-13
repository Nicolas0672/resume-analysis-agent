from io import BytesIO

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph

def open_docx(file_bytes: bytes) -> Document:
    try:
        return Document(BytesIO(file_bytes))
    except (PackageNotFoundError, Exception):
        raise ValueError("Invalid DOCX file")

def parse_docx(doc: Document) -> list[dict[str, str]]:
    paragraphs = []
    sentence_id = 0

    def process_paragraph(p: Paragraph):
        nonlocal sentence_id
        text = p.text.strip()
        if text:
            paragraphs.append({
                "sentence_id": f"{sentence_id}",
                "text": text,
            })
            sentence_id += 1

    def process_table(table: Table):
        seen_cells = set()
        for row in table.rows:
            for cell in row.cells:
                if cell._tc in seen_cells:
                    continue
                seen_cells.add(cell._tc)
                for child in cell._tc:
                    if isinstance(child, CT_P):
                        process_paragraph(Paragraph(child, doc))
                    elif isinstance(child, CT_Tbl):
                        process_table(Table(child, doc))

    def iter_blocks(parent):
        for child in parent:
            if isinstance(child, CT_P):
                process_paragraph(Paragraph(child, doc))
            elif isinstance(child, CT_Tbl):
                process_table(Table(child, doc))
            elif child.tag.endswith("sdt"):
                sdt_content = child.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}sdtContent")
                if sdt_content is not None:
                    iter_blocks(sdt_content)

    iter_blocks(doc.element.body)

    # Fallback to standard collections if custom iteration found nothing but items exist
    if not paragraphs and (doc.paragraphs or doc.tables):
        for p in doc.paragraphs:
            process_paragraph(p)
        for t in doc.tables:
            process_table(t)

    return paragraphs

# Alias parse_doc to parse_docx for compatibility
parse_doc = parse_docx

def parse_resume(file_bytes: bytes):
    doc = open_docx(file_bytes)
    return parse_docx(doc)




