import io
import re
from typing import Optional
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor

from model.job_pydantic import (
    ResumeEducation,
    ResumeExperience,
    ResumeLeadership,
    ResumeProject,
    ResumeSkills,
    ResumeStructure,
)

# Standard ATS color palette matching web preview
COLOR_TEXT_PRIMARY = RGBColor(17, 24, 39)    # #111827 - dark zinc/black
COLOR_TEXT_MUTED = RGBColor(107, 114, 128)   # #6B7280 - muted zinc
COLOR_BORDER = "CCCCCC"                       # #CCCCCC - light grey rule
FONT_NAME = "Arial"


def _set_cell_margins(cell, top=0, bottom=0, left=0, right=0):
    """Sets explicit 0 margins/padding for table cells to guarantee tight alignment."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for m, val in [("top", top), ("bottom", bottom), ("left", left), ("right", right)]:
        node = OxmlElement(f"w:{m}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def _add_section_heading(doc: Document, title: str):
    """Adds a section heading with an uppercase title, tight spacing, and a bottom border rule."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True

    run = p.add_run(title.upper())
    run.bold = True
    run.font.name = FONT_NAME
    run.font.size = Pt(10.5)
    run.font.color.rgb = COLOR_TEXT_PRIMARY

    # Inject XML bottom border (matching web preview's border-b)
    pPr = p._p.get_or_add_pPr()
    pBdr = parse_xml(
        f'<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f'<w:bottom w:val="single" w:sz="6" w:space="1" w:color="{COLOR_BORDER}"/>'
        f"</w:pBdr>"
    )
    pPr.append(pBdr)
    return p


def _add_two_column_row(
    doc: Document,
    left_text: str,
    right_text: str = "",
    left_bold: bool = False,
    left_italic: bool = False,
    right_bold: bool = False,
    right_muted: bool = True,
    space_before: float = 1.0,
    space_after: float = 0.5,
    font_size: float = 10.0,
):
    """Creates a borderless 2-column tabular row for robust ATS left/right alignment.
    Total width: 7.5 inches (printable area for 8.5" width with 0.5" margins).
    """
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Set column widths
    table.columns[0].width = Inches(5.4)
    table.columns[1].width = Inches(2.1)

    # Strip table borders
    tblPr = table._tbl.tblPr
    tblBorders = parse_xml(
        '<w:tblBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:top w:val="none"/><w:left w:val="none"/><w:bottom w:val="none"/>'
        '<w:right w:val="none"/><w:insideH w:val="none"/><w:insideV w:val="none"/>'
        "</w:tblBorders>"
    )
    tblPr.append(tblBorders)

    # Left cell
    cell_l = table.cell(0, 0)
    cell_l.width = Inches(5.4)
    _set_cell_margins(cell_l, top=0, bottom=0, left=0, right=0)
    p_l = cell_l.paragraphs[0]
    p_l.paragraph_format.space_before = Pt(space_before)
    p_l.paragraph_format.space_after = Pt(space_after)
    p_l.paragraph_format.line_spacing = 1.05

    run_l = p_l.add_run(left_text)
    run_l.bold = left_bold
    run_l.italic = left_italic
    run_l.font.name = FONT_NAME
    run_l.font.size = Pt(font_size)
    run_l.font.color.rgb = COLOR_TEXT_PRIMARY

    # Right cell
    cell_r = table.cell(0, 1)
    cell_r.width = Inches(2.1)
    _set_cell_margins(cell_r, top=0, bottom=0, left=0, right=0)
    p_r = cell_r.paragraphs[0]
    p_r.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_r.paragraph_format.space_before = Pt(space_before)
    p_r.paragraph_format.space_after = Pt(space_after)
    p_r.paragraph_format.line_spacing = 1.05

    if right_text:
        run_r = p_r.add_run(right_text)
        run_r.bold = right_bold
        run_r.font.name = FONT_NAME
        run_r.font.size = Pt(font_size - 0.5)
        run_r.font.color.rgb = COLOR_TEXT_MUTED if right_muted else COLOR_TEXT_PRIMARY

    return table


def _add_bullet_item(doc: Document, text: str):
    """Adds a single bullet point item using Word's native List Bullet style."""
    clean_text = re.sub(r"^[\s\-•*·–—]+\s*", "", str(text or "").strip())
    if not clean_text:
        return
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(1.5)
    p.paragraph_format.line_spacing = 1.05
    p.paragraph_format.left_indent = Inches(0.22)

    run = p.add_run(clean_text)
    run.font.name = FONT_NAME
    run.font.size = Pt(9.5)
    run.font.color.rgb = COLOR_TEXT_PRIMARY
    return p


def _extract_active_skill_entries(skills: Optional[ResumeSkills]) -> list[tuple[str, list[str]]]:
    """Extracts non-empty skills categories with human-readable labels matching utils.ts."""
    if not skills:
        return []

    category_map = [
        ("programming_languages", "Languages"),
        ("frameworks", "Frameworks"),
        ("libraries", "Libraries"),
        ("databases", "Databases"),
        ("cloud", "Cloud / DevOps"),
        ("tools", "Developer Tools"),
        ("other", "Other"),
    ]

    entries = []
    for attr, label in category_map:
        items = getattr(skills, attr, None)
        if isinstance(items, list):
            filtered = [str(x).strip() for x in items if str(x).strip()]
            if filtered:
                entries.append((label, filtered))
    return entries


def generate_docx(resume: ResumeStructure) -> io.BytesIO:
    """Generates an ATS-compliant, high-fidelity DOCX document matching the tailored resume preview.

    Layout Specifications:
    - Page Setup: US Letter (8.5" x 11.0") with 0.5" margins on all four sides.
    - Standard Typography: Arial (18pt Name, 9.5pt Contact, 10.5pt Section Headings, 9.5-10pt Body).
    - Section Sequence: Header -> Education -> Work Experience -> Projects -> Leadership -> Skills.
    - Two-column alignment for entry headers (Title/Institution on left, Date/Location on right).
    - Underline border rule beneath each uppercase section heading.
    - Native Word List Bullet points with compact ATS spacing.
    """
    doc = Document()

    # Configure 0.5-inch page margins (Letter size: 8.5 x 11 inches)
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11.0)
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)

    # 1. Header: Candidate Name
    p_name = doc.add_paragraph()
    p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_name.paragraph_format.space_before = Pt(0)
    p_name.paragraph_format.space_after = Pt(1)

    candidate_name = (resume.name or "Candidate Name").strip()
    r_name = p_name.add_run(candidate_name.upper())
    r_name.bold = True
    r_name.font.name = FONT_NAME
    r_name.font.size = Pt(18)
    r_name.font.color.rgb = COLOR_TEXT_PRIMARY

    # 2. Header: Contact Line
    if resume.contact:
        contact_str = resume.contact.strip()
        if contact_str:
            p_contact = doc.add_paragraph()
            p_contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_contact.paragraph_format.space_before = Pt(0)
            p_contact.paragraph_format.space_after = Pt(4)

            r_contact = p_contact.add_run(contact_str)
            r_contact.font.name = FONT_NAME
            r_contact.font.size = Pt(9.5)
            r_contact.font.color.rgb = COLOR_TEXT_MUTED

    # 3. Education Section (At top, matching web preview)
    if resume.education and len(resume.education) > 0:
        _add_section_heading(doc, "Education")
        for edu in resume.education:
            # Row 1: Institution (bold, left) | Duration or Location (right)
            inst_text = (edu.institution or "Institution").strip()
            date_loc = (edu.duration or edu.location or "").strip()
            _add_two_column_row(
                doc,
                left_text=inst_text,
                right_text=date_loc,
                left_bold=True,
                space_before=2.0,
                space_after=0.5,
                font_size=10.0,
            )

            # Row 2: Degree in Field of Study (left) | GPA (right)
            degree_parts = []
            if edu.degree:
                degree_parts.append(edu.degree.strip())
            if edu.field_of_study:
                degree_parts.append(f"in {edu.field_of_study.strip()}")
            degree_str = " ".join(degree_parts) or "Degree"

            gpa_str = f"GPA: {edu.gpa.strip()}" if edu.gpa else ""
            _add_two_column_row(
                doc,
                left_text=degree_str,
                right_text=gpa_str,
                left_bold=False,
                left_italic=False,
                space_before=0.5,
                space_after=0.5,
                font_size=9.5,
            )

            # Row 3: Coursework (if available)
            if edu.coursework and len(edu.coursework) > 0:
                clean_courses = [str(c).strip() for c in edu.coursework if str(c).strip()]
                if clean_courses:
                    p_course = doc.add_paragraph()
                    p_course.paragraph_format.space_before = Pt(0.5)
                    p_course.paragraph_format.space_after = Pt(2.0)
                    p_course.paragraph_format.left_indent = Inches(0.0)

                    r_c_label = p_course.add_run("Coursework: ")
                    r_c_label.font.name = FONT_NAME
                    r_c_label.font.size = Pt(9.0)
                    r_c_label.font.color.rgb = COLOR_TEXT_MUTED

                    r_c_val = p_course.add_run(", ".join(clean_courses))
                    r_c_val.font.name = FONT_NAME
                    r_c_val.font.size = Pt(9.0)
                    r_c_val.font.color.rgb = COLOR_TEXT_MUTED

    # 4. Work Experience Section
    if resume.work_experience and len(resume.work_experience) > 0:
        _add_section_heading(doc, "Work Experience")
        for exp in resume.work_experience:
            # Row 1: Job Title (bold, left) | Duration (right)
            _add_two_column_row(
                doc,
                left_text=(exp.job_title or "Position Title").strip(),
                right_text=(exp.duration or "").strip(),
                left_bold=True,
                space_before=3.0,
                space_after=0.5,
                font_size=10.0,
            )

            # Row 2: Company • Location (left)
            comp_parts = []
            if exp.company:
                comp_parts.append(exp.company.strip())
            if exp.location:
                comp_parts.append(f"• {exp.location.strip()}")
            comp_str = " ".join(comp_parts) if comp_parts else ""

            if comp_str:
                p_comp = doc.add_paragraph()
                p_comp.paragraph_format.space_before = Pt(0)
                p_comp.paragraph_format.space_after = Pt(1.5)
                r_comp = p_comp.add_run(comp_str)
                r_comp.font.name = FONT_NAME
                r_comp.font.size = Pt(9.5)
                r_comp.font.color.rgb = COLOR_TEXT_MUTED

            # Bullets
            for bullet in exp.bullets:
                _add_bullet_item(doc, bullet.text)

    # 5. Projects Section
    if resume.projects and len(resume.projects) > 0:
        _add_section_heading(doc, "Projects")
        for proj in resume.projects:
            # Row 1: Project Name (bold, left) | Duration (right)
            _add_two_column_row(
                doc,
                left_text=(proj.project_name or "Project Name").strip(),
                right_text=(proj.duration or "").strip(),
                left_bold=True,
                space_before=3.0,
                space_after=0.5,
                font_size=10.0,
            )

            # Row 2: Technologies (left)
            if proj.technologies and len(proj.technologies) > 0:
                tech_str = ", ".join([str(t).strip() for t in proj.technologies if str(t).strip()])
                if tech_str:
                    p_tech = doc.add_paragraph()
                    p_tech.paragraph_format.space_before = Pt(0)
                    p_tech.paragraph_format.space_after = Pt(1.5)

                    r_t_label = p_tech.add_run("Technologies: ")
                    r_t_label.font.name = FONT_NAME
                    r_t_label.font.size = Pt(9.0)
                    r_t_label.font.color.rgb = COLOR_TEXT_MUTED

                    r_t_val = p_tech.add_run(tech_str)
                    r_t_val.font.name = FONT_NAME
                    r_t_val.font.size = Pt(9.0)
                    r_t_val.font.color.rgb = COLOR_TEXT_MUTED

            # Bullets
            for bullet in proj.bullets:
                _add_bullet_item(doc, bullet.text)

    # 6. Leadership & Extracurriculars Section
    if resume.leadership and len(resume.leadership) > 0:
        _add_section_heading(doc, "Leadership & Extracurriculars")
        for lead in resume.leadership:
            # Row 1: Title (bold, left) | Position or Duration (right)
            right_lead = (lead.position or lead.duration or "").strip()
            _add_two_column_row(
                doc,
                left_text=(lead.title or "Organization / Activity").strip(),
                right_text=right_lead,
                left_bold=True,
                space_before=3.0,
                space_after=1.5,
                font_size=10.0,
            )

            # Bullets
            for bullet in lead.bullets:
                _add_bullet_item(doc, bullet.text)

    # 7. Technical Skills Section
    skill_entries = _extract_active_skill_entries(resume.skills)
    if skill_entries:
        _add_section_heading(doc, "Technical Skills")
        for label, items in skill_entries:
            p_skill = doc.add_paragraph()
            p_skill.paragraph_format.space_before = Pt(1.0)
            p_skill.paragraph_format.space_after = Pt(1.5)
            p_skill.paragraph_format.line_spacing = 1.05

            r_label = p_skill.add_run(f"{label}: ")
            r_label.bold = True
            r_label.font.name = FONT_NAME
            r_label.font.size = Pt(9.5)
            r_label.font.color.rgb = COLOR_TEXT_PRIMARY

            r_items = p_skill.add_run(", ".join(items))
            r_items.font.name = FONT_NAME
            r_items.font.size = Pt(9.5)
            r_items.font.color.rgb = COLOR_TEXT_PRIMARY

    # Stream out document buffer
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer
