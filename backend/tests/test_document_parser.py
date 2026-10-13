from io import BytesIO
import unittest
import docx

from services.document_parser import (
    open_docx,
    parse_doc,
    parse_docx,
    parse_resume,
)


class TestDocumentParser(unittest.TestCase):
    def test_parse_docx_paragraphs_only(self):
        doc = docx.Document()
        doc.add_paragraph("Jane Doe")
        doc.add_paragraph("")  # Empty paragraph should be skipped
        doc.add_paragraph("Software Engineer")

        results = parse_docx(doc)

        self.assertEqual(len(results), 2)
        self.assertEqual(results[0], {"sentence_id": "0", "text": "Jane Doe"})
        self.assertEqual(results[1], {"sentence_id": "1", "text": "Software Engineer"})

    def test_parse_docx_with_table(self):
        doc = docx.Document()
        doc.add_paragraph("Technical Skills")

        tbl = doc.add_table(rows=2, cols=2)
        tbl.cell(0, 0).text = "Languages"
        tbl.cell(0, 1).text = "Python, TypeScript, SQL"
        tbl.cell(1, 0).text = "Databases"
        tbl.cell(1, 1).text = "PostgreSQL, Redis"

        results = parse_docx(doc)

        expected = [
            {"sentence_id": "0", "text": "Technical Skills"},
            {"sentence_id": "1", "text": "Languages"},
            {"sentence_id": "2", "text": "Python, TypeScript, SQL"},
            {"sentence_id": "3", "text": "Databases"},
            {"sentence_id": "4", "text": "PostgreSQL, Redis"},
        ]
        self.assertEqual(results, expected)

    def test_parse_docx_document_order(self):
        doc = docx.Document()
        doc.add_paragraph("Header Before Table")

        tbl = doc.add_table(rows=1, cols=2)
        tbl.cell(0, 0).text = "Table Left"
        tbl.cell(0, 1).text = "Table Right"

        doc.add_paragraph("Footer After Table")

        results = parse_docx(doc)

        self.assertEqual(len(results), 4)
        self.assertEqual(results[0]["text"], "Header Before Table")
        self.assertEqual(results[1]["text"], "Table Left")
        self.assertEqual(results[2]["text"], "Table Right")
        self.assertEqual(results[3]["text"], "Footer After Table")

    def test_parse_docx_merged_cells_deduplication(self):
        doc = docx.Document()
        tbl = doc.add_table(rows=2, cols=2)
        # Row 0 merged header
        tbl.cell(0, 0).text = "Merged Section Header"
        tbl.cell(0, 0).merge(tbl.cell(0, 1))

        # Row 1 individual columns
        tbl.cell(1, 0).text = "Col A"
        tbl.cell(1, 1).text = "Col B"

        results = parse_docx(doc)

        texts = [item["text"] for item in results]
        # "Merged Section Header" should only appear once, not twice
        self.assertEqual(texts, ["Merged Section Header", "Col A", "Col B"])

    def test_parse_docx_multi_paragraph_cells(self):
        doc = docx.Document()
        tbl = doc.add_table(rows=1, cols=1)
        cell = tbl.cell(0, 0)
        cell.text = "First bullet in cell."
        cell.add_paragraph("Second bullet in cell.")

        results = parse_docx(doc)

        self.assertEqual(len(results), 2)
        self.assertEqual(results[0], {"sentence_id": "0", "text": "First bullet in cell."})
        self.assertEqual(results[1], {"sentence_id": "1", "text": "Second bullet in cell."})

    def test_parse_docx_empty_and_whitespace_cells(self):
        doc = docx.Document()
        tbl = doc.add_table(rows=2, cols=2)
        tbl.cell(0, 0).text = "   "
        tbl.cell(0, 1).text = ""
        tbl.cell(1, 0).text = "Valid Cell Text"
        tbl.cell(1, 1).text = "   \n\t  "

        results = parse_docx(doc)

        self.assertEqual(len(results), 1)
        self.assertEqual(results[0], {"sentence_id": "0", "text": "Valid Cell Text"})

    def test_parse_doc_alias(self):
        doc = docx.Document()
        doc.add_paragraph("Testing Alias")

        results = parse_doc(doc)
        self.assertEqual(results, [{"sentence_id": "0", "text": "Testing Alias"}])

    def test_parse_resume_end_to_end(self):
        doc = docx.Document()
        doc.add_paragraph("Jane Doe")
        tbl = doc.add_table(rows=1, cols=2)
        tbl.cell(0, 0).text = "Skills"
        tbl.cell(0, 1).text = "Python, FastAPI"

        buf = BytesIO()
        doc.save(buf)
        file_bytes = buf.getvalue()

        results = parse_resume(file_bytes)

        self.assertEqual(len(results), 3)
        self.assertEqual(results[0]["text"], "Jane Doe")
        self.assertEqual(results[1]["text"], "Skills")
        self.assertEqual(results[2]["text"], "Python, FastAPI")

    def test_open_docx_invalid_bytes(self):
        with self.assertRaises(ValueError):
            open_docx(b"corrupted or invalid docx bytes")


if __name__ == "__main__":
    unittest.main()
