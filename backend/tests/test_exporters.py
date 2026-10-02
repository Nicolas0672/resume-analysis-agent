import io
import pytest
from docx import Document

from backend.model.job_pydantic import (
    ResumeBullet,
    ResumeEducation,
    ResumeExperience,
    ResumeLeadership,
    ResumeProject,
    ResumeSkills,
    ResumeStructure,
)
from backend.services.docx_exporter import generate_docx
from backend.services.pdf_exporter import generate_pdf


def sample_resume():
    return ResumeStructure(
        name="Alex Mercer",
        contact="alex.mercer@email.com | (555) 234-5678 | github.com/alexmercer | San Francisco, CA",
        education=[
            ResumeEducation(
                institution="University of California, Berkeley",
                degree="Bachelor of Science",
                field_of_study="Computer Science",
                duration="2020 - 2024",
                gpa="3.85",
                coursework=["Algorithms & Data Structures", "Operating Systems", "Distributed Systems"],
            )
        ],
        work_experience=[
            ResumeExperience(
                job_title="Software Engineering Intern",
                company="Tech Innovations Corp",
                location="San Francisco, CA",
                duration="June 2023 - Aug 2023",
                bullets=[
                    ResumeBullet(text="Engineered microservices using Python & FastAPI, reducing API latency by 35%.", sentence_id=1),
                    ResumeBullet(text="Automated CI/CD pipelines with GitHub Actions, accelerating release cycles.", sentence_id=2),
                ],
            )
        ],
        projects=[
            ResumeProject(
                project_name="Distributed Key-Value Store",
                duration="Spring 2024",
                technologies=["Go", "Raft", "gRPC", "Docker"],
                bullets=[
                    ResumeBullet(text="Implemented Raft consensus algorithm handling leader election & log replication.", sentence_id=3),
                ],
            )
        ],
        leadership=[
            ResumeLeadership(
                title="Computer Science Undergraduate Association",
                position="President",
                duration="2023 - 2024",
                bullets=[
                    ResumeBullet(text="Organized technical hackathons for over 500+ participants with 15 corporate sponsors.", sentence_id=4),
                ],
            )
        ],
        skills=ResumeSkills(
            programming_languages=["Python", "Go", "TypeScript", "SQL"],
            frameworks=["FastAPI", "React", "Next.js"],
            databases=["PostgreSQL", "Redis"],
            cloud=["AWS", "Docker", "Kubernetes"],
            tools=["Git", "Linux"],
            other=["GraphQL"],
        ),
    )


def test_generate_docx_success():
    resume = sample_resume()
    buf = generate_docx(resume)

    assert isinstance(buf, io.BytesIO)
    content = buf.getvalue()
    assert len(content) > 0

    # Ensure python-docx can reopen and parse the generated docx file
    doc = Document(buf)
    all_text = "\n".join([p.text for p in doc.paragraphs])
    assert "ALEX MERCER" in all_text
    assert "alex.mercer@email.com" in all_text
    assert "EDUCATION" in all_text
    assert "WORK EXPERIENCE" in all_text
    assert "PROJECTS" in all_text
    assert "LEADERSHIP & EXTRACURRICULARS" in all_text
    assert "TECHNICAL SKILLS" in all_text
    assert "Languages: Python, Go, TypeScript, SQL" in all_text


def test_generate_pdf_success():
    resume = sample_resume()
    buf = generate_pdf(resume)

    assert isinstance(buf, io.BytesIO)
    content = buf.getvalue()
    assert len(content) > 0
    # PDF magic byte signature
    assert content.startswith(b"%PDF-")


def test_exporters_with_minimal_resume():
    # Tests minimal resume with missing optional fields to prevent NoneType or KeyError crashes
    minimal = ResumeStructure(
        name="Jane Doe",
        contact=None,
        education=[],
        work_experience=[],
        projects=[],
        leadership=[],
        skills=None,
    )

    docx_buf = generate_docx(minimal)
    assert len(docx_buf.getvalue()) > 0
    doc = Document(docx_buf)
    assert any("JANE DOE" in p.text for p in doc.paragraphs)

    pdf_buf = generate_pdf(minimal)
    assert pdf_buf.getvalue().startswith(b"%PDF-")


def test_pdf_xml_escaping():
    # Verify characters like &, <, > in bullet or project text don't break ReportLab XML parser
    resume = ResumeStructure(
        name="Dev & Tester <Lead>",
        contact="dev@test.com & co",
        education=[],
        work_experience=[
            ResumeExperience(
                job_title="R&D Engineer <L4>",
                company="Bits & Bytes Corp",
                bullets=[
                    ResumeBullet(text="Managed CI/CD <pipelines> & deployed to A/B test groups (x < 10 && y > 20).", sentence_id=1),
                ],
            )
        ],
        projects=[],
        leadership=[],
        skills=ResumeSkills(programming_languages=["C++", "C#", "HTML & CSS"]),
    )

    pdf_buf = generate_pdf(resume)
    assert pdf_buf.getvalue().startswith(b"%PDF-")

    docx_buf = generate_docx(resume)
    assert len(docx_buf.getvalue()) > 0


def test_pdf_bullets_not_literal_bull():
    resume = sample_resume()
    pdf_buf = generate_pdf(resume)
    raw_content = pdf_buf.getvalue()

    # The PDF output must NEVER contain literal "&bull;" string
    assert b"&bull;" not in raw_content
    # Leading dash / bullet in source text should be cleanly stripped
    dash_resume = ResumeStructure(
        name="Alex Mercer",
        education=[],
        work_experience=[
            ResumeExperience(
                job_title="Software Engineer",
                company="Tech Co",
                bullets=[ResumeBullet(text="- Engineered microservices", sentence_id=1)],
            )
        ],
        projects=[],
        leadership=[],
        skills=None,
    )
    pdf_dash = generate_pdf(dash_resume)
    assert b"&bull;" not in pdf_dash.getvalue()
